import random
import time
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from statistics import fmean

import pandas as pd
from django.core.management.base import BaseCommand
from openpyxl.styles import PatternFill

from Clients.models import Client
from Dishes.models import Dish
from Menus.generator.domain import DietConfig, MealSlotConfig, ROLE_LABELS, day_micro_totals, day_totals
from Menus.generator.fitness import fitness
from Menus.generator.meal_structure import get_meal_structure
from Menus.generator.micronutrients import MICRO_IDS, ranges_for_client
from Menus.generator.pools import build_candidate_pools
from Menus.generator.service import GenerationError, _generate_day_within_ranges
from Menus.generator.targets import attach_micro_ranges, macro_split_for_kcal
from Products.models import ProductMicronutrient
from Users.models import User

KCALS = [1500, 2000, 2500, 3000, 3500]
SIZES = ['S', 'M', 'L', 'XL']

INGESTA_SETS = [
    ('lunch_dinner', ['breakfast', 'lunch', 'dinner']),
    ('five_slot', ['breakfast', 'mid_morning_snack', 'lunch', 'afternoon_snack', 'dinner']),
]

LEGACY_SINGLE_TO_INTAKE_NAME = {
    'breakfast': 'Desayuno',
    'mid_morning_snack': 'Media mañana',
    'afternoon_snack': 'Merienda',
}

VIT_A_ID, PHOSPHORUS_ID = 8, 28
INFO_ONLY_IDS = {'vit_a': VIT_A_ID, 'fosforo': PHOSPHORUS_ID}
ALL_MICRO_IDS = {**MICRO_IDS, **INFO_ONLY_IDS}

DAYS_PER_WEEK = 7


def _load_extra_micros(dish_ids, cache):
    """Menus.generator.pools._product_micro_map solo trae los MICRO_IDS que usa
    el generador (11), asi que DishCandidate.micros_100g nunca lleva vitamina A
    ni fosforo. Para poder reportarlos como informativos hay que recalcularlos
    aparte, con la misma formula que Menus.generator.pools._dish_micros_100g,
    pero filtrando solo a INFO_ONLY_IDS."""
    dishes = Dish.objects.filter(id__in=dish_ids).prefetch_related('dishproduct_set__product')
    product_ids = {dp.product_id for dish in dishes for dp in dish.dishproduct_set.all()}

    product_micro_map = defaultdict(dict)
    for pm in ProductMicronutrient.objects.filter(
        product_id__in=product_ids, micronutrient_id__in=INFO_ONLY_IDS.values()
    ):
        product_micro_map[pm.product_id][pm.micronutrient_id] = float(pm.value)

    for dish in dishes:
        total_quantity = 0
        totals = defaultdict(float)
        for dish_product in dish.dishproduct_set.all():
            quantity = dish_product.quantity or 0
            total_quantity += quantity
            for micro_id, value in product_micro_map.get(dish_product.product_id, {}).items():
                totals[micro_id] += value * quantity / 100
        cache[dish.id] = (
            {micro_id: value / total_quantity * 100 for micro_id, value in totals.items()}
            if total_quantity else {}
        )


ABOVE_MAX_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')  # rojo
BELOW_MIN_FILL = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')  # azul


def highlight_out_of_range(worksheet, combos_df, ranges):
    """Colorea la celda avg_<micro> de cada fila cuya media semanal cae fuera de
    ranges_for_client: rojo si supera el maximo, azul si no llega al minimo."""
    if combos_df.empty:
        return

    col_index = {name: idx + 1 for idx, name in enumerate(combos_df.columns)}
    for name, mid in MICRO_IDS.items():
        avg_col = col_index.get(f'avg_{name}')
        if avg_col is None:
            continue
        lo, hi = ranges.get(mid, (None, None))
        for offset, value in enumerate(combos_df[f'avg_{name}']):
            if hi is not None and value > hi:
                worksheet.cell(row=offset + 2, column=avg_col).fill = ABOVE_MAX_FILL
            elif lo is not None and value < lo:
                worksheet.cell(row=offset + 2, column=avg_col).fill = BELOW_MIN_FILL


def day_extra_micro_totals(day_menu, extra_cache):
    totals = {}
    for courses in day_menu.values():
        for entry in courses.values():
            if entry is None:
                continue
            candidate, qty = entry
            factor = qty / 100.0
            for micro_id, value in extra_cache.get(candidate.dish_id, {}).items():
                totals[micro_id] = totals.get(micro_id, 0.0) + value * factor
    return totals


def build_meal_slots(ingesta_names, dessert_flag, starter_flag, standalone_by_name, groups):
    slots = []
    for name in ingesta_names:
        if name in LEGACY_SINGLE_TO_INTAKE_NAME:
            intake = standalone_by_name[LEGACY_SINGLE_TO_INTAKE_NAME[name]]
            slots.append(MealSlotConfig(
                key=f'single-{intake.id}', label=intake.name, kind='single',
                intakes={'single': intake},
            ))
        elif name == 'lunch':
            slots.append(MealSlotConfig(
                key='group-comida', label='Comida', kind='group', intakes=groups['comida'],
                include_starter=bool(starter_flag), include_dessert=bool(dessert_flag),
            ))
        elif name == 'dinner':
            slots.append(MealSlotConfig(
                key='group-cena', label='Cena', kind='group', intakes=groups['cena'],
                include_starter=bool(starter_flag), include_dessert=bool(dessert_flag),
            ))
        else:
            raise ValueError(f'Ingesta legacy desconocida: {name!r}')
    return slots


def iter_combos(limit=None):
    n = 0
    for kcal in KCALS:
        for size in SIZES:
            for ingesta_key, ingesta_names in INGESTA_SETS:
                for starter_flag in (0, 1):
                    for dessert_flag in (0, 1):
                        if limit is not None and n >= limit:
                            return
                        yield n, kcal, size, ingesta_key, ingesta_names, dessert_flag, starter_flag
                        n += 1


def run_combo(client, nutri, pool_cache, kcal, size, ingesta_key, ingesta_names,
              dessert_flag, starter_flag, standalone_by_name, groups, rng):
    meal_slots = build_meal_slots(ingesta_names, dessert_flag, starter_flag, standalone_by_name, groups)

    pool_key = (ingesta_key, dessert_flag, starter_flag, size)
    pools = pool_cache.get(pool_key)
    if pools is None:
        pools = build_candidate_pools(client, nutri, meal_slots, portion_size=size)
        pool_cache[pool_key] = pools

    for slot in meal_slots:
        for role in slot.active_roles():
            if not pools[slot.key].get(role):
                raise GenerationError(
                    f'No hay platos de tipo «{ROLE_LABELS[role]}» activos y compatibles '
                    f'con este cliente para «{slot.label}».'
                )

    target = macro_split_for_kcal(kcal)
    attach_micro_ranges(target, client)
    config = DietConfig(
        days=DAYS_PER_WEEK, start_date=date(2026, 1, 5), meal_slots=meal_slots,
        target=target, portion_size=size,
    )

    day_menus = [_generate_day_within_ranges(pools, config, rng) for _ in range(config.days)]
    return day_menus, target


def aggregate_week(day_menus, target, extra_micro_cache):
    dish_ids = {
        candidate.dish_id
        for day_menu in day_menus
        for courses in day_menu.values()
        for entry in courses.values()
        if entry is not None
        for candidate, _qty in [entry]
    }
    missing = [dish_id for dish_id in dish_ids if dish_id not in extra_micro_cache]
    if missing:
        _load_extra_micros(missing, extra_micro_cache)

    kcal_l, prot_l, fat_l, carb_l, fit_l = [], [], [], [], []
    micro_l = {mid: [] for mid in ALL_MICRO_IDS.values()}

    for day_menu in day_menus:
        k, p, f, c = day_totals(day_menu)
        kcal_l.append(k)
        prot_l.append(p)
        fat_l.append(f)
        carb_l.append(c)
        fit_l.append(fitness(day_menu, target))
        micros = day_micro_totals(day_menu)
        extra_micros = day_extra_micro_totals(day_menu, extra_micro_cache)
        for mid in MICRO_IDS.values():
            micro_l[mid].append(micros.get(mid, 0.0))
        for mid in INFO_ONLY_IDS.values():
            micro_l[mid].append(extra_micros.get(mid, 0.0))

    avg_kcal, avg_prot, avg_fat, avg_carb = fmean(kcal_l), fmean(prot_l), fmean(fat_l), fmean(carb_l)

    def absol_relativo(grams, kcal_per_g):
        contrib = grams * kcal_per_g
        return contrib / target.kcal, contrib / avg_kcal

    prot_absol, prot_rel = absol_relativo(avg_prot, 4)
    fat_absol, fat_rel = absol_relativo(avg_fat, 9)
    carb_absol, carb_rel = absol_relativo(avg_carb, 4)

    return {
        'avg_fitness': fmean(fit_l),
        'desv_kcal_dia': avg_kcal - target.kcal,
        'avg_kcal': avg_kcal,
        'prot_absol': prot_absol, 'prot_relativo': prot_rel, 'avg_prot_g': avg_prot,
        'fat_absol': fat_absol, 'fat_relativo': fat_rel, 'avg_fat_g': avg_fat,
        'carb_absol': carb_absol, 'carb_relativo': carb_rel, 'avg_carb_g': avg_carb,
        **{f'avg_{name}': fmean(micro_l[mid]) for name, mid in ALL_MICRO_IDS.items()},
    }


class Command(BaseCommand):
    help = (
        'Replica el test por lotes del generador legacy (micros_mujer.ods): barre 160 '
        'combinaciones de kcal/tamaño/ingestas/postre/entrante, genera una dieta de 7 días '
        'para cada una con el generador actual, promedia la semana y compara los límites '
        'de Menus.generator.micronutrients.ranges_for_client con la media obtenida.'
    )

    def add_arguments(self, parser):
        parser.add_argument('--seed', type=int, default=42)
        parser.add_argument('--limit', type=int, default=None)
        parser.add_argument('--output', type=str, default=None)
        parser.add_argument('--nutritionist-email', type=str, default='c_soraya5@hotmail.com')
        parser.add_argument('--client-dni', type=str, default='MICROTEST-0001')

    def handle(self, *args, **options):
        from django.conf import settings

        nutri = User.objects.get(email=options['nutritionist_email'])

        client, created = Client.objects.get_or_create(
            user=nutri, dni=options['client_dni'],
            defaults=dict(
                email='micro-week-test-client@idiet.local',
                first_name='Test', last_name='MicroWeekWoman',
                birth_date=date(1994, 6, 15), gender='Female',
                height=165, weight=Decimal('60.00'),
                phone_number='600000000', phone_number_2='',
                address='N/A', postal_code='28080', city='Madrid',
                activity_level='Moderada', status=True,
            ),
        )
        self.stdout.write(
            self.style.SUCCESS(f'Cliente de test id={client.id} ({"creado" if created else "reutilizado"})')
        )

        standalone, groups = get_meal_structure()
        standalone_by_name = {intake.name: intake for intake in standalone}
        ranges = ranges_for_client(client)

        rng = random.Random(options['seed'])
        pool_cache = {}
        extra_micro_cache = {}
        rows, failures = [], []
        combos = list(iter_combos(options['limit']))
        total = len(combos)
        t0 = time.monotonic()

        for idx, kcal, size, ingesta_key, ingesta_names, dessert_flag, starter_flag in combos:
            try:
                day_menus, target = run_combo(
                    client, nutri, pool_cache, kcal, size, ingesta_key, ingesta_names,
                    dessert_flag, starter_flag, standalone_by_name, groups, rng,
                )
            except GenerationError as exc:
                failures.append({
                    'combo_index': idx, 'kcal': kcal, 'size': size,
                    'ingestas': str(ingesta_names), 'dessert': dessert_flag,
                    'starter': starter_flag, 'error': str(exc),
                })
                self.stdout.write(self.style.WARNING(f'[{idx + 1}/{total}] SALTADA: {exc}'))
                continue

            stats = aggregate_week(day_menus, target, extra_micro_cache)
            row = {
                'combo_index': idx, 'kcal_target': kcal, 'size': size,
                'ingestas': str(ingesta_names), 'dessert': dessert_flag, 'starter': starter_flag,
                **stats,
                **{
                    f'exceeds_{name}': (
                        (ranges.get(mid, (None, None))[0] is not None
                         and stats[f'avg_{name}'] < ranges[mid][0])
                        or (ranges.get(mid, (None, None))[1] is not None
                            and stats[f'avg_{name}'] > ranges[mid][1])
                    )
                    for name, mid in MICRO_IDS.items()
                },
            }
            rows.append(row)

            if (idx + 1) % 10 == 0 or (idx + 1) == total:
                elapsed = time.monotonic() - t0
                self.stdout.write(
                    f'[{idx + 1}/{total}] elapsed={elapsed:.1f}s ok={len(rows)} failed={len(failures)}'
                )

        elapsed_total = time.monotonic() - t0

        summary_rows = []
        n_ok = len(rows)
        for name, mid in MICRO_IDS.items():
            lo, hi = ranges.get(mid, (None, None))
            n_exceeding = sum(1 for r in rows if r[f'exceeds_{name}']) if n_ok else 0
            summary_rows.append({
                'micronutriente': name, 'micronutriente_id': mid,
                'min_bound': lo, 'max_bound': hi,
                'n_evaluados': n_ok, 'n_exceden': n_exceeding,
                'pct_exceden': (n_exceeding / n_ok) if n_ok else None,
            })

        meta = {
            'generated_at': datetime.now().isoformat(timespec='seconds'),
            'seed': options['seed'],
            'nutritionist_email': options['nutritionist_email'],
            'client_id': client.id,
            'total_combos_attempted': total,
            'n_ok': n_ok,
            'n_failed': len(failures),
            'elapsed_seconds': round(elapsed_total, 1),
            'macro_split': 'sistema actual: 20% prot / 25% fat / 55% carb (legacy usaba 15/30/55)',
            'avg_fitness_note': (
                'avg_fitness usa Menus.generator.fitness.fitness (pesos/objetivo del sistema actual); '
                'no es comparable numéricamente con el avg_fitness del fichero legacy.'
            ),
        }

        output_path = options['output'] or str(settings.BASE_DIR / 'micros_mujer_actual.xlsx')
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            combos_df = pd.DataFrame(rows)
            combos_df.to_excel(writer, sheet_name='Combos', index=False)
            pd.DataFrame(summary_rows).to_excel(writer, sheet_name='Summary', index=False)
            pd.DataFrame([meta]).to_excel(writer, sheet_name='Meta', index=False)
            if failures:
                pd.DataFrame(failures).to_excel(writer, sheet_name='Failures', index=False)

            highlight_out_of_range(writer.sheets['Combos'], combos_df, ranges)

        self.stdout.write(self.style.SUCCESS(
            f'Listo: {n_ok}/{total} combinaciones generadas, {len(failures)} fallidas, '
            f'{elapsed_total:.1f}s -> {output_path}'
        ))
